from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import PatternFill

from llm_prep.ingest.xlsx_ingest import ingest_xlsx


def test_xlsx_ingest_extracts_table_and_colors(tmp_path: Path) -> None:
    xlsx_path = tmp_path / "sample.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Tier", "Value"])
    ws.append(["A", 100])
    ws.append(["B", 200])
    ws["A2"].fill = PatternFill(fill_type="solid", fgColor="FFFF0000")
    wb.save(xlsx_path)

    parsed = ingest_xlsx(xlsx_path, {"csv_sample_rows": 10})
    assert parsed["tables"]
    table = parsed["tables"][0]
    assert table["table_kind"] == "xlsx_sheet_table"
    assert table["metadata"]["sheet_name"] == "Sheet1"
    assert table["metadata"]["color_counts"]

