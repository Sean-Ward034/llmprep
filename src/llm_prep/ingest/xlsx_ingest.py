"""XLSX ingestion."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from llm_prep.extract.tables import infer_schema

PARSER_NAME = "xlsx_ingest"
PARSER_VERSION = "1.0.0"


def _is_row_empty(row_values: list[Any]) -> bool:
    return all(cell in (None, "") for cell in row_values)


def _color_token(cell) -> str | None:
    fill = getattr(cell, "fill", None)
    if not fill:
        return None
    fg = getattr(fill, "fgColor", None)
    if not fg:
        return None
    if fg.type == "rgb" and fg.rgb:
        return str(fg.rgb).upper()
    if fg.type == "indexed":
        return f"UNKNOWN:indexed:{fg.indexed}"
    if fg.type == "theme":
        return f"UNKNOWN:theme:{fg.theme}"
    return None


def ingest_xlsx(path: Path, config: dict) -> dict:
    warnings: list[str] = []
    tables: list[dict] = []
    wb = load_workbook(path, data_only=True)

    for sheet in wb.worksheets:
        all_rows = list(sheet.iter_rows())
        if not all_rows:
            continue

        header_idx = None
        for idx, row in enumerate(all_rows):
            values = [cell.value for cell in row]
            if not _is_row_empty(values):
                header_idx = idx
                break
        if header_idx is None:
            continue

        header_row = all_rows[header_idx]
        headers = [(str(cell.value).strip() if cell.value is not None else "") for cell in header_row]
        headers = [h or f"column_{i+1}" for i, h in enumerate(headers)]

        data_rows: list[list[Any]] = []
        color_counts: dict[str, int] = {}
        for row in all_rows[header_idx + 1 :]:
            values = [cell.value for cell in row]
            if _is_row_empty(values):
                continue
            data_rows.append(values)
            for cell in row:
                token = _color_token(cell)
                if not token:
                    continue
                color_counts[token] = color_counts.get(token, 0) + 1

        if not data_rows:
            continue

        max_col = len(headers)
        min_row = header_idx + 1
        max_row = header_idx + 1 + len(data_rows)
        range_hint = f"A{min_row}:{get_column_letter(max_col)}{max_row}"
        schema = infer_schema(headers, data_rows[: int(config.get("csv_sample_rows", 50))])
        if any(key.startswith("UNKNOWN:") for key in color_counts):
            warnings.append(f"Unknown color encoding in {path}:{sheet.title}")

        tables.append(
            {
                "table_kind": "xlsx_sheet_table",
                "headers": headers,
                "rows": data_rows,
                "schema": schema.model_dump(),
                "metadata": {
                    "sheet_name": sheet.title,
                    "range": range_hint,
                    "color_counts": color_counts,
                },
                "locator": f"xlsx:{sheet.title}!{range_hint}",
            }
        )
    return {
        "parser": {"name": PARSER_NAME, "version": PARSER_VERSION},
        "text_segments": [],
        "tables": tables,
        "warnings": warnings,
    }

